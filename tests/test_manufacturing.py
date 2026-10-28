from pathlib import Path

import openpyxl
import pytest

from cad_inventory.manufacturing.readers import read_workbook, dictionary_from_workbook
from cad_inventory.manufacturing.processes import normalized_id, parse_time, extract_operations, build_routes, build_electrode_graph
from cad_inventory.manufacturing.matching import match_family


def operation(part, row, code, time):
    return {'mold_id':'SMX0100','part_id':part,'part_id_raw':part,'part_name_raw':'电极' if '=' in part else '塑模镶件',
            'source_file':'source.xls','sheet':'模具进度查询','row_number':row,'record_id':f'source.xls::模具进度查询::R{row}',
            'process_code':code,'time_numeric':float(time),'time_decimal':str(time),'time_raw':str(time),
            'time_nature':'quota; actual_processing_time_unconfirmed','process_dictionary_matches':[],
            'raw_fields':{'件号':part,'工序名称':code,'工时定额':str(time)}}


def test_strict_join_preserves_asterisk_and_electrode_suffix():
    assert normalized_id(' smx0100 - m20a* ')=='SMX0100-M20A*'
    assert normalized_id('SMX0100-F20A=1=3')=='SMX0100-F20A=1=3'
    routes=build_routes([operation('SMX0100-M20A*',3,'MS2',1),operation('SMX0100-M20A*=1=3',4,'CNC3',2)])
    cad={'drawings':[{'mold_id':'SMX0100','drawing_number':'SMX0100-M20A','sources':[]}]}
    result=match_family('SMX0100',cad,routes)
    assert result['strict_matched_primary_count']==0
    assert len(result['unconfirmed_alias_candidates'])==1
    assert not result['electrode_parent_links'][0]['parent_exact_cad_match']
    assert result['unmatched_excel_electrode_ids']==['SMX0100-M20A*=1=3']


def test_repeated_operations_preserved_and_sorted_by_original_row():
    part='SMX0100-F20A=1=3'
    records=[operation(part,103,'EDM2',3.5),operation(part,101,'CNC3',1.5),operation(part,102,'CNC3',1.5)]
    route=build_routes(records)[0]
    assert route['source_rows']==[101,102,103]
    assert route['operation_codes']==['CNC3','CNC3','EDM2']
    assert route['numeric_time_sum']==6.5
    assert route['repeated_identical_operation_extra_occurrences']==1
    assert len({o['record_id'] for o in route['operations']})==3


def test_electrode_dependency_ledger_no_double_count_and_cost_unconfirmed():
    parent='SMX0100-F20A'
    records=[operation(parent,28,'MS2',4),operation(parent,29,'TG2',2.5),operation(parent,30,'EDM1',4),operation(parent,31,'EDM2',0),
             operation(parent+'=1=3',101,'CNC3',1.5),operation(parent+'=1=3',102,'CNC3',1.5),operation(parent+'=1=3',103,'EDM2',3.5),
             operation(parent+'=2=3',104,'CNC3',0.7),operation(parent+'=2=3',105,'EDM2',2)]
    graph=build_electrode_graph(build_routes(records),parent,ppt_confirmed=True)
    assert graph['body_numeric_sum']==10.5
    assert graph['electrode_manufacture_numeric_sum']==3.7
    assert graph['electrode_discharge_numeric_sum']==5.5
    assert graph['distinct_source_record_numeric_sum']==19.7
    assert graph['source_record_count']==9 and graph['double_counted_source_records']==0
    assert all(row['cost_owner']=='unconfirmed' for row in graph['ledger'])
    assert len([e for e in graph['edges'] if e['relation']=='discharge_detail_of'])==2
    assert len([e for e in graph['edges'] if e['relation']=='electrode_ready_for_EDM2'])==2
    assert records[3]['time_raw']=='0'


def test_units_and_invalid_times_never_imputed():
    assert parse_time('')[1]=='missing'
    assert parse_time('not available')[1]=='non_numeric'
    assert parse_time('NaN')[1]=='non_finite'
    assert parse_time('-1')[1]=='negative'
    assert parse_time('0')[1]=='zero'


def test_dictionary_only_resolves_explicit_merged_categories(tmp_path):
    path=tmp_path/'codes.xlsx'
    workbook=openpyxl.Workbook();s=workbook.active
    s.append(['说明'])
    s.cell(4,1,'加工中心');s.cell(4,2,'CNC');s.cell(4,3,'CNC3');s.cell(4,4,'铜电极')
    s.cell(5,3,'CNC5');s.cell(5,4,'硬料精加工')
    s.merge_cells('A4:A5');s.merge_cells('B4:B5')
    s.cell(6,3,'RCL');s.cell(6,4,'热处理')
    workbook.save(path)
    result=dictionary_from_workbook(read_workbook(path))
    entries={e['code']:e for e in result['entries']}
    assert entries['CNC5']['category']=='加工中心'
    assert entries['CNC5']['category_evidence'][0]['source_row']==4
    assert entries['CNC5']['raw_values'][0] is None
    assert entries['RCL']['category'] is None
    assert entries['RCL']['description_raw']=='热处理'


def test_workbook_row_provenance_and_time_label(tmp_path):
    path=tmp_path/'process.xlsx';book=openpyxl.Workbook();s=book.active
    s.append(['模具内部编号','件号','零件名称','工序名称','工时定额'])
    s.append(['SMX0100','SMX0100-X01A','塑模镶件','TG2','0'])
    book.save(path)
    records,profiles=extract_operations(read_workbook(path),{'entries':[]})
    assert records[0]['row_number']==2
    assert records[0]['raw_fields']['工时定额']=='0'
    assert records[0]['raw_cells'][4]['coordinate']=='E2'
    assert not records[0]['full_table_unit_confirmed']
    assert records[0]['time_nature']=='quota; actual_processing_time_unconfirmed'
    assert profiles[0]['column_storage_types']['工时定额']=={'text':1}


def test_actual_legacy_workbooks_and_dictionary_regression():
    root=Path(__file__).resolve().parents[1]
    source=root/'data/raw/manufacturing'
    if not (source/'SMX0100-详细工艺工时.xls').exists():
        pytest.skip('Customer records not present')
    dictionary=dictionary_from_workbook(read_workbook(source/'模具加工工艺分类.xlsx'))
    assert dictionary['entry_count']==37
    lookup={e['code']:e for e in dictionary['entries']}
    assert lookup['CNC3']['description_raw']=='铜电极'
    assert lookup['RCL']['category'] is None
    for family,count,parts,zeros in [('CMY1000A',783,188,9),('SMX0100',151,44,9)]:
        workbook=read_workbook(source/f'{family}-详细工艺工时.xls')
        assert workbook['reader']=='xlrd'
        records,_=extract_operations(workbook,dictionary)
        assert len(records)==count
        assert len({r['part_id'] for r in records})==parts
        assert sum(r['time_status']=='zero' for r in records)==zeros
        assert all(r['process_mapping_status']=='mapped' for r in records)
        if family=='SMX0100':
            graph=build_electrode_graph(build_routes(records),'SMX0100-F20A',ppt_confirmed=True)
            assert graph['distinct_source_record_numeric_sum']==19.7
            assert [r['row_number'] for r in graph['ledger']]==[28,29,30,31,101,102,103,104,105]


def test_explicit_frame_part_number_does_not_expand_or_truncate_suffix():
    row=operation('CMY1000A-P20A-B',3,'EW3',0.7)
    row['mold_id']='CMY1000A'
    routes=build_routes([row])
    cad={'drawings':[{'mold_id':'CMY1000A','drawing_number':'CMY1000A-P20A',
                      'cad_part_ids':['CMY1000A-P20A-B'],'sources':[
                          {'source_file':'drawing.dxf','part_number_evidence':[{'raw_part_number':'P20A-B','attribute_tag':'PARTNO','frame_insert_handle':'123'}]}]}]}
    result=match_family('CMY1000A',cad,routes)
    assert result['strict_matched_primary_count']==0
    assert result['evidence_matched_primary_count']==1
    assert result['cad_unique_part_number_count']==1
    match=result['part_number_evidence_matches'][0]
    assert match['part_id']=='CMY1000A-P20A-B'
    assert match['cad_drawing_numbers']==['CMY1000A-P20A']
    assert not match['version_alignment_confirmed']


def test_incomplete_electrode_is_flagged_without_imputing_manufacture():
    parent='SMX0100-M24B'
    records=[operation(parent,82,'TG2',1),operation(parent,83,'EDM2',0),
             operation(parent+'=1=3',153,'EDM2',2)]
    graph=build_electrode_graph(build_routes(records),parent)
    assert not graph['branch_completeness'][0]['has_manufacturing_record']
    assert graph['branch_completeness'][0]['has_discharge_record']
    assert graph['electrode_manufacture_numeric_sum']==0
    assert graph['source_record_count']==3
    assert graph['status']=='candidate; engineer confirmation needed'


def test_title_detail_conflict_is_not_treated_as_clean_match():
    row=operation('CMY1000A-D21A',10,'TG2',1)
    row['mold_id']='CMY1000A'
    cad={'drawings':[{'mold_id':'CMY1000A','drawing_number':'CMY1000A-D21A',
                      'cad_part_ids':['CMY1000A-D21AG'],'sources':[]}]}
    result=match_family('CMY1000A',cad,build_routes([row]))
    assert result['strict_matched_primary_count']==1
    assert result['evidence_matched_primary_count']==0
    assert result['combined_evidence_primary_count']==1
    assert result['combined_evidence_no_identity_conflict_count']==0
    assert result['identifier_conflicts'][0]['detail_part_ids']==['CMY1000A-D21AG']
