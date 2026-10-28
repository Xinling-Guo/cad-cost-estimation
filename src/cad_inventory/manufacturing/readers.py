"""Read legacy XLS, XLSX merged-cell evidence, and complete PPTX structure."""
from collections import Counter
from pathlib import Path

import openpyxl
from pptx import Presentation
import xlrd


def read_workbook(path):
    path = Path(path)
    sheets = []
    if path.suffix.lower() == '.xls':
        book = xlrd.open_workbook(path, formatting_info=True)
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                cells = [{'column': c + 1, 'coordinate': xlrd.cellname(r, c),
                          'raw_value': sheet.cell_value(r, c),
                          'storage_type': {0:'empty',1:'text',2:'number',3:'date',4:'boolean',5:'error',6:'blank'}.get(sheet.cell_type(r,c), 'unknown')}
                         for c in range(sheet.ncols)]
                rows.append({'row_number': r + 1, 'cells': cells})
            sheets.append({'sheet': sheet.name, 'row_count': sheet.nrows, 'column_count': sheet.ncols,
                           'merged_ranges': [[a+1,b,c+1,d] for a,b,c,d in sheet.merged_cells], 'rows': rows})
        book.release_resources()
        reader = 'xlrd'
    else:
        book = openpyxl.load_workbook(path, data_only=False)
        cache = openpyxl.load_workbook(path, data_only=True)
        dimensions = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for sheet in book:
            rows = []
            # Retain declared trailing empty rows (the source taxonomy includes row41).
            declared = dimensions[sheet.title]
            declared_extent = declared.calculate_dimension()
            max_row, max_column = max(sheet.max_row,declared.max_row), max(sheet.max_column,declared.max_column)
            for row in sheet.iter_rows(min_row=1,max_row=max_row,min_col=1,max_col=max_column):
                cells = [{'column': c.column, 'coordinate': c.coordinate, 'raw_value': c.value,
                          'storage_type': 'formula' if c.data_type == 'f' else 'empty' if c.value is None else 'text' if isinstance(c.value,str) else 'number',
                          'cached_value': cache[sheet.title].cell(c.row,c.column).value if c.data_type == 'f' else None}
                         for c in row]
                rows.append({'row_number': row[0].row, 'cells': cells})
            sheets.append({'sheet': sheet.title, 'row_count': sheet.max_row, 'column_count': sheet.max_column,
                           'declared_dimension':declared_extent,
                           'merged_ranges': [str(r) for r in sheet.merged_cells.ranges], 'rows': rows})
        book.close();cache.close();dimensions.close();reader = 'openpyxl'
    return {'source_file': path.as_posix(), 'reader': reader, 'sheets': sheets}


def values(row):
    return [cell['raw_value'] for cell in row['cells']]


def dictionary_from_workbook(workbook):
    """Resolve only explicitly merged category cells; never forward-fill blanks."""
    entries = []
    for sheet in workbook['sheets']:
        anchors = {}
        for region in sheet['merged_ranges']:
            if not isinstance(region,str):
                continue
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            c0,r0,c1,r1 = range_boundaries(region)
            for r in range(r0,r1+1):
                for c in range(c0,c1+1):
                    anchors[(r,c)] = (r0,c0)
        for row in sheet['rows']:
            raw = values(row)
            if row['row_number'] < 4 or len(raw)<4 or not raw[2] or not raw[3]:
                continue
            resolved, evidence = [], []
            for c in (1,2):
                r0,c0 = anchors.get((row['row_number'],c),(row['row_number'],c))
                val = sheet['rows'][r0-1]['cells'][c0-1]['raw_value']
                resolved.append(val)
                evidence.append({'column': c, 'source_row':r0, 'source_column':c0, 'value':val,
                                 'resolution':'merged_cell_anchor' if (r0,c0)!=(row['row_number'],c) else 'original_cell'})
            entries.append({'source_file':workbook['source_file'], 'sheet':sheet['sheet'], 'row_number':row['row_number'],
                            'code':str(raw[2]).strip(), 'description_raw':raw[3], 'category':resolved[0], 'category_code':resolved[1],
                            'category_evidence':evidence,'raw_values':raw})
    duplicates = {k:v for k,v in Counter(e['code'] for e in entries).items() if v>1}
    return {'entries':entries,'entry_count':len(entries),'duplicate_codes':duplicates,
            'missing_category_codes':[e['code'] for e in entries if not e['category'] or not e['category_code']]}


def read_presentation(path, image_dir):
    presentation = Presentation(path)
    slides = []
    image_dir.mkdir(parents=True, exist_ok=True)
    def shapes_content(shapes, slide_number):
        result = []
        for shape in shapes:
            item = {'shape_id':shape.shape_id,'name':shape.name,'shape_type':int(shape.shape_type),
                    'position_emu':[shape.left,shape.top,shape.width,shape.height]}
            if shape.has_text_frame:
                item['text'] = shape.text
            if shape.has_table:
                item['table'] = [[c.text for c in row.cells] for row in shape.table.rows]
            if shape.shape_type == 6:
                item['children'] = shapes_content(shape.shapes,slide_number)
            if shape.shape_type == 13:
                filename=f'slide_{slide_number}_shape_{shape.shape_id}.{shape.image.ext}'
                (image_dir/filename).write_bytes(shape.image.blob)
                item['image_file']=(image_dir/filename).as_posix()
            starts=shape.element.xpath('.//a:stCxn'); ends=shape.element.xpath('.//a:endCxn')
            if starts or ends:
                item['connector']={'start_shape_id':int(starts[0].get('id')) if starts else None,
                                   'end_shape_id':int(ends[0].get('id')) if ends else None,
                                   'head_end_types':[n.get('type') for n in shape.element.xpath('.//a:headEnd')],
                                   'tail_end_types':[n.get('type') for n in shape.element.xpath('.//a:tailEnd')]}
            result.append(item)
        return result
    for number,slide in enumerate(presentation.slides,1):
        slides.append({'slide_number':number,'shapes':shapes_content(slide.shapes,number),
                       'notes':slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None})
    return {'source_file':str(path),'slide_count':len(slides),'slides':slides,
            'verified_business_rules':[
                {'slide':3,'rule':'CMY1000A 为冲压模具，仅2D；SMX0100 为注塑模具，说明称2D均有3D。'},
                {'slide':4,'rule':'图号对应Excel件号，同件号工序按从上至下排列。'},
                {'slide':5,'rule':'F20A 主体 MS2→TG2→EDM1→EDM2；两个铜电极支线汇入EDM2。'},
                {'slide':5,'rule':'电极1 CNC3 1.5h→CNC3 1.5h，放电3.5h；电极2 CNC3 0.7h，放电2.0h。'}],
            'unit_evidence':{'slide':5,'unit':'h','scope':'F20A example','full_workbook_unit_confirmed':False},
            'actual_time_confirmed':False}
